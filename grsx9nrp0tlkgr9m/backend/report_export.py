from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def write_markdown(path: Path, markdown_text: str) -> Path:
    path.write_text(markdown_text, encoding="utf-8")
    return path


def write_docx(path: Path, markdown_text: str) -> Path:
    doc = Document()
    for line in markdown_text.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        else:
            doc.add_paragraph(line)
    doc.save(path)
    return path


def write_pdf(path: Path, markdown_text: str) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 40
    for raw in markdown_text.splitlines():
        line = raw[:110]
        c.drawString(40, y, line)
        y -= 14
        if y < 40:
            c.showPage()
            y = height - 40
    c.save()
    return path
